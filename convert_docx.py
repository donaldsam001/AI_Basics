import os
import pandas as pd
from docx import Document

# =========================
# CONFIG
# =========================
INPUT_CSV = "dataset.csv"
OUTPUT_DIR = "example_data/docx"

# =========================
# CREATE OUTPUT DIRECTORY
# =========================
os.makedirs(OUTPUT_DIR, exist_ok=True)

# =========================
# READ CSV
# =========================
df = pd.read_csv(INPUT_CSV)

# =========================
# CONVERT EACH ROW -> DOCX
# =========================
for index, row in df.iterrows():

    # Use resume_id as filename if available
    resume_id = str(row.get("resume_id", f"resume_{index}"))

    # Remove invalid characters from filename
    safe_name = "".join(
        c for c in resume_id
        if c.isalnum() or c in ("_", "-", ".")
    )

    if not safe_name:
        safe_name = f"resume_{index}"

    output_path = os.path.join(
        OUTPUT_DIR,
        f"{safe_name}.docx"
    )

    # Create Word document
    doc = Document()

    # Title
    doc.add_heading(f"Resume {resume_id}", level=1)

    # Add every column
    for column in df.columns:

        value = row[column]

        # Skip empty values
        if pd.isna(value):
            continue

        value = str(value).strip()

        if not value:
            continue

        # Field name
        doc.add_heading(column.replace("_", " ").title(), level=2)

        # Field content
        doc.add_paragraph(value)

    # Save
    doc.save(output_path)

    print(f"Created: {output_path}")

print("\nDone!")
print(f"Total resumes: {len(df)}")
print(f"Output directory: {OUTPUT_DIR}")